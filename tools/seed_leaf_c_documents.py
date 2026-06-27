import html
import json
import mimetypes
import re
import secrets
import sys
from http.cookiejar import CookieJar
from pathlib import Path
from urllib.parse import urlencode
from urllib.request import HTTPCookieProcessor, Request, build_opener

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_URL = "http://127.0.0.1:18088"
DATA_DIR = Path(r"C:\Users\DELL\Desktop\Data")
USERNAME = "admin"
PASSWORD = "admin"

HEADERS = [
    "STT",
    "Họ và tên",
    "Sinh ngày",
    "Trình độ",
    "Cấp bậc",
    "Mức độ hoàn thành nhiệm vụ",
    "Hoạt động của đơn vị",
]
FIRST_NAMES = ["Nguyễn Văn", "Trần Minh", "Lê Hoàng", "Phạm Quang", "Đỗ Anh", "Vũ Thành", "Hoàng Đức", "Bùi Hữu", "Đặng Quốc", "Phan Hải"]
EDUCATIONS = ["THPT", "Đại học", "Thạc sĩ", "Tiến sĩ"]
RANKS = ["Cấp tá", "Cấp úy", "Hạ sỹ quan", "Dân sự"]
COMPLETIONS = ["Xuất sắc", "Tốt", "Hoàn thành", "Không hoàn thành"]
ACTIVITY_POOL = [
    "Diễn tập phòng chống thiên tai",
    "Sinh hoạt chuyên đề chuyển đổi số",
    "Tổ chức hội thao cấp đơn vị",
    "Kết nghĩa với địa phương",
    "Kiểm tra bảo đảm an toàn thông tin",
    "Bồi dưỡng nghiệp vụ văn thư lưu trữ",
    "Tổng vệ sinh doanh trại cuối tuần",
    "Tập huấn sơ cấp cứu tại chỗ",
    "Rà soát hồ sơ quản lý cán bộ",
    "Hội nghị rút kinh nghiệm quý",
    "Giao lưu văn nghệ chào mừng ngày truyền thống",
    "Tuyên truyền pháp luật cho đoàn viên",
    "Kiểm tra nền nếp trực ban",
    "Tổ chức ngày kỹ thuật cấp đơn vị",
    "Phối hợp tuần tra địa bàn",
    "Bồi dưỡng kỹ năng làm việc nhóm",
]


def activities_for_unit(unit_number):
    start = (unit_number * 3) % len(ACTIVITY_POOL)
    return [ACTIVITY_POOL[(start + offset) % len(ACTIVITY_POOL)] for offset in range(4)]


def request(opener, url, data=None, headers=None):
    req = Request(url, data=data, headers=headers or {})
    return opener.open(req, timeout=60)


def login_and_dashboard():
    opener = build_opener(HTTPCookieProcessor(CookieJar()))
    payload = urlencode({"username": USERNAME, "password": PASSWORD}).encode()
    request(
        opener,
        f"{BASE_URL}/login",
        payload,
        {"Content-Type": "application/x-www-form-urlencoded;charset=UTF-8"},
    ).read()
    dashboard = request(opener, f"{BASE_URL}/").read().decode("utf-8", errors="replace")
    csrf_match = re.search(r'data-tree-csrf="([^"]*)"', dashboard)
    data_match = re.search(
        r'<script id="dashboard-report-data" type="application/json">(.*?)</script>',
        dashboard,
        re.S,
    )
    if not csrf_match or not data_match:
        raise RuntimeError("Could not read dashboard csrf/report data after login")
    csrf = html.unescape(csrf_match.group(1))
    units = json.loads(html.unescape(data_match.group(1)))
    return opener, csrf, units


def rows_for_unit(unit_name):
    number_match = re.search(r"\d+", unit_name)
    unit_number = int(number_match.group(0)) if number_match else 1
    unit_activities = activities_for_unit(unit_number)
    rows = [HEADERS]
    for index in range(1, 21):
        seed = unit_number + index
        birth_year = 1972 + (seed * 3) % 31
        birth_month = 1 + (seed % 12)
        birth_day = 1 + ((seed * 2) % 27)
        display_name = FIRST_NAMES[seed % len(FIRST_NAMES)]
        rows.append([
            index,
            display_name,
            f"{birth_day:02d}/{birth_month:02d}/{birth_year}",
            EDUCATIONS[seed % len(EDUCATIONS)],
            RANKS[(seed + 1) % len(RANKS)],
            COMPLETIONS[(seed + 2) % len(COMPLETIONS)],
            "\n".join(unit_activities) if index == 1 else "",
        ])
    return rows


def write_workbook(unit_name):
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    path = DATA_DIR / f"{unit_name}.xlxs"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = unit_name[:31]
    for row in rows_for_unit(unit_name):
        sheet.append(row)
    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for column in sheet.columns:
        max_len = max(len(str(cell.value or "")) for cell in column)
        sheet.column_dimensions[get_column_letter(column[0].column)].width = min(max(max_len + 3, 10), 42)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    workbook.save(path)
    return path


def multipart_form(fields, file_field, file_path):
    boundary = f"----WebsiteBuu{secrets.token_hex(12)}"
    chunks = []
    for name, value in fields.items():
        chunks.append(f"--{boundary}\r\n".encode())
        chunks.append(f'Content-Disposition: form-data; name="{name}"\r\n\r\n'.encode())
        chunks.append(str(value).encode("utf-8"))
        chunks.append(b"\r\n")
    mime = mimetypes.guess_type(str(file_path))[0] or "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    chunks.append(f"--{boundary}\r\n".encode())
    chunks.append(
        f'Content-Disposition: form-data; name="{file_field}"; filename="{file_path.name}"\r\n'.encode()
    )
    chunks.append(f"Content-Type: {mime}\r\n\r\n".encode())
    chunks.append(file_path.read_bytes())
    chunks.append(b"\r\n")
    chunks.append(f"--{boundary}--\r\n".encode())
    return boundary, b"".join(chunks)


def upload_document(opener, csrf, unit, path):
    fields = {
        "csrf": csrf,
        "org_id": unit["id"],
        "year": "2026",
        "return_to": "/",
        "title": unit["name"],
    }
    boundary, body = multipart_form(fields, "document", path)
    response = request(
        opener,
        f"{BASE_URL}/documents",
        body,
        {
            "Content-Type": f"multipart/form-data; boundary={boundary}",
            "Accept": "text/html,application/xhtml+xml",
        },
    )
    return response.status


def main():
    opener, csrf, units = login_and_dashboard()
    c_units = sorted(
        (unit for unit in units if re.fullmatch(r"c\d+", unit.get("name", ""), re.I)),
        key=lambda unit: int(re.search(r"\d+", unit["name"]).group(0)),
    )
    if len(c_units) != 72:
        raise RuntimeError(f"Expected 72 c units, found {len(c_units)}")
    uploaded = []
    for unit in c_units:
        path = write_workbook(unit["name"])
        status = upload_document(opener, csrf, unit, path)
        uploaded.append({"name": unit["name"], "id": unit["id"], "file": str(path), "status": status})
        print(f"uploaded {unit['name']} -> {status}")
    print(json.dumps({"uploaded": uploaded}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(f"ERROR: {error}", file=sys.stderr)
        raise
